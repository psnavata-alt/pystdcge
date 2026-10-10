"""World oil-price shocks with sourced shares (added in 0.3).

With the petroleum-split SAM (paths.SAM_2023_PETR; tools/build_sam_petroleum_split.py) a world oil-price change enters the
CGE where it happens:
  refined petroleum imports (cpetr)   world import price x (1 + oil change)  [product prices assumed to move one-for-one
                                       with crude; Dubai and Singapore product prices co-move]
  crude imports (inside mining, cmine) world import price x (1 + oil change x crude share of mining imports)
The combined 2023 SAM (published results) keeps the old approximation: cchem import price x (1 + oil x 0.6), 0.6 being an
assumed petroleum share of cchem imports. PSA 2018 puts that share at 0.512 (imports column, sectors 107-108 of 107-118).

For the I-O price model the refined-petroleum price change is set directly (IOTable.price). Its pass-through from a world
oil-price change, from PSA's 80-industry table (industry 35, coke and refined petroleum; industry 21, crude petroleum and
natural gas) and the 2018 Supply table (imports of refined petroleum, sector 107):
    p = s_m + (1 - s_m) (c + o p)   ->   p = (s_m + (1 - s_m) c) / (1 - (1 - s_m) o)
s_m = imported share of refined-petroleum supply; c = crude input per peso of refinery output; o = refined petroleum
used per peso of refinery output (blending and own use, whose price also rises). Replaces the assumed 0.6.
"""
from __future__ import annotations

import numpy as np
import openpyxl

from . import paths

PSA_240 = paths.DATA_RAW / "io" / "240-Industry20Table202018.xlsx"
PSA_SUT = paths.DATA_RAW / "io" / "2018 Supply and Use Tables.xlsx"


def _imports_240() -> np.ndarray:
    ws = openpyxl.load_workbook(PSA_240, read_only=True, data_only=True)["240 TT"]
    rows = list(ws.iter_rows(values_only=True))
    j = [str(h).strip() if h else "" for h in rows[9]].index("Imports of Goods and Services")
    return np.array([float(rows[9 + i][j] or 0) for i in range(1, 241)])


# The three shares below, computed by import_shares() from PSA's 2018 240-industry input-output table (imports
# column). Stored because the PSA table may not be redistributed; recomputed when you have downloaded it.
IMPORT_SHARES_PSA2018 = {"crude_in_mining": 0.6765750743551517, "fertilizer_in_chemicals": 0.025993988387202443,
                         "petroleum_in_chem_and_petroleum": 0.5122653869965557}


def import_shares() -> dict:
    """Shares in PSA 2018 imports (240-industry table, imports column)."""
    if not PSA_240.exists():
        return dict(IMPORT_SHARES_PSA2018)
    m = _imports_240()
    mining = m[38:45].sum()                       # sectors 39-45
    chem = m[108:118].sum()                       # sectors 109-118
    return {"crude_in_mining": float(m[39] / mining),                 # sector 40 crude petroleum and natural gas
            "fertilizer_in_chemicals": float(m[109] / chem),          # sector 110 fertilizers and nitrogen compounds
            "petroleum_in_chem_and_petroleum": float(m[106:108].sum() / (m[106:108].sum() + chem))}


def io_refined_passthrough(io=None) -> dict:
    """Pass-through of a world oil-price change to the price of refined petroleum in the I-O price model."""
    from .iomodel import IOTable
    io = io or IOTable()
    k = next(i for i, n in enumerate(io.names) if n.startswith("Coke and refined petroleum"))
    cr = next(i for i, n in enumerate(io.names) if n.startswith("Crude petroleum"))
    X = io.Z.sum(axis=0) if not hasattr(io, "X") else io.X
    c = float(io.Z[cr, k] / X[k])
    o = float(io.Z[k, k] / X[k])
    ws = openpyxl.load_workbook(PSA_SUT, read_only=True, data_only=True)["240_Supply"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() if h else "" for h in rows[9]]
    jd, jm = hdr.index("Total Domestic Production"), hdr.index("Total Imports")
    r107 = next(r for r in rows if r[3] is not None and str(r[3]).strip() == "107")
    s_m = float(r107[jm]) / (float(r107[jd]) + float(r107[jm]))
    p = (s_m + (1 - s_m) * c) / (1 - (1 - s_m) * o)
    return {"passthrough": p, "import_share": s_m, "crude_coef": c, "own_use_coef": o}


def cge_oil_shock(m, oil: float, fertilizer: float = 0.0, combined_share: float = 0.6) -> dict:
    """CGE shock spec for a world oil-price change `oil` (0.40 = +40%) and optionally a fertilizer price change.
    Uses the split SAM's accounts when the model has 'cpetr'; otherwise the old combined approximation."""
    sh = import_shares()
    pwm = {}
    if "cpetr" in m.C:
        pwm["cpetr"] = 1 + oil
        pwm["cmine"] = 1 + oil * sh["crude_in_mining"]
        if fertilizer:
            pwm["cchem"] = 1 + fertilizer * sh["fertilizer_in_chemicals"]
    else:
        pwm["cchem"] = 1 + oil * combined_share
    return {"pwm_scale": pwm}
